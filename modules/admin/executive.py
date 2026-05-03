from __future__ import annotations

import datetime
from flask import jsonify, request, render_template, send_file

from db import get_conn
from services.brand import get_brand


def _parse_date(s: str | None, default: datetime.date) -> datetime.date:
    if not s:
        return default
    s = (s or "").strip()
    try:
        return datetime.date.fromisoformat(s[:10])
    except Exception:
        return default


def _date_range_from_request() -> tuple[datetime.date, datetime.date]:
    today = datetime.date.today()
    d_to = _parse_date(request.args.get("to"), today)
    d_from = _parse_date(request.args.get("from"), d_to - datetime.timedelta(days=30))
    if d_from > d_to:
        d_from, d_to = d_to, d_from
    return d_from, d_to


def _station_filter_clause(station_id: int | None, field: str = "station_id") -> tuple[str, list]:
    if station_id is None:
        return "", []
    return f" AND {field}=?", [int(station_id)]


def _compute_summary(brand: str, station_id: int | None, d_from: datetime.date, d_to: datetime.date) -> dict:
    conn = get_conn()
    cur = conn.cursor()

    st = None
    if station_id is not None:
        cur.execute("SELECT id, code, name FROM stations WHERE brand=? AND id=?", (brand, int(station_id)))
        r = cur.fetchone()
        if r:
            st = dict(r)

    rng_sql = "date(created_at) BETWEEN date(?) AND date(?)"
    params_rng = [d_from.isoformat(), d_to.isoformat()]

    # Alerts
    sfx, sparams = _station_filter_clause(station_id)
    cur.execute(f"SELECT COUNT(1) AS c FROM alerts WHERE brand=? AND status='open'{sfx}", (brand, *sparams))
    alerts_open = int(cur.fetchone()["c"] or 0)
    cur.execute(f"SELECT COUNT(1) AS c FROM alerts WHERE brand=? AND {rng_sql}{sfx}", (brand, *params_rng, *sparams))
    alerts_created = int(cur.fetchone()["c"] or 0)

    # Maintenance
    cur.execute(f"SELECT COUNT(1) AS c FROM maintenance WHERE brand=? AND {rng_sql}{sfx}", (brand, *params_rng, *sparams))
    maint_created = int(cur.fetchone()["c"] or 0)

    # Pipas
    cur.execute(f"SELECT COUNT(1) AS c FROM pipas WHERE brand=? AND {rng_sql}{sfx}", (brand, *params_rng, *sparams))
    pipas_created = int(cur.fetchone()["c"] or 0)

    # Payments
    cur.execute(f"SELECT COUNT(1) AS c FROM payments WHERE brand=? AND status='pending'{sfx}", (brand, *sparams))
    payments_pending = int(cur.fetchone()["c"] or 0)
    cur.execute(f"SELECT COUNT(1) AS c FROM payments WHERE brand=? AND status='validated'{sfx}", (brand, *sparams))
    payments_validated = int(cur.fetchone()["c"] or 0)

    # Activities (Consulting only) - submissions in range
    cur.execute(
        f"SELECT status, COUNT(1) AS c FROM submissions WHERE brand=? AND {rng_sql}{sfx} GROUP BY status",
        (brand, *params_rng, *sparams),
    )
    sub_by_status = {r["status"]: int(r["c"] or 0) for r in cur.fetchall()}

    # Events planned in range (calendar)
    # For station-specific: events where station_id is NULL or == station
    if station_id is None:
        cur.execute(
            "SELECT COUNT(1) AS c FROM calendar_events WHERE brand=? AND date(start_date) BETWEEN date(?) AND date(?)",
            (brand, d_from.isoformat(), d_to.isoformat()),
        )
        events_planned = int(cur.fetchone()["c"] or 0)
    else:
        cur.execute(
            "SELECT COUNT(1) AS c FROM calendar_events WHERE brand=? AND date(start_date) BETWEEN date(?) AND date(?) AND (station_id IS NULL OR station_id=?)",
            (brand, d_from.isoformat(), d_to.isoformat(), int(station_id)),
        )
        events_planned = int(cur.fetchone()["c"] or 0)

    # Documental reviews (SASISOPA / SGM)
    # join requirements to filter station
    doc_station_clause = ""
    doc_station_params: list = []
    if station_id is not None:
        doc_station_clause = " AND (r.station_id IS NULL OR r.station_id=?)"
        doc_station_params = [int(station_id)]
    cur.execute(
        """
        SELECT s.module, COUNT(1) AS c
        FROM doc_submissions s
        JOIN doc_requirements r ON r.id=s.requirement_id AND r.brand=s.brand
        WHERE s.brand=? AND s.review_status='PENDING'""" + doc_station_clause + " GROUP BY s.module",
        (brand, *doc_station_params),
    )
    pending_by_module = { (r["module"] or "").lower(): int(r["c"] or 0) for r in cur.fetchall() }
    sasisopa_pending = int(pending_by_module.get("sasisopa", 0))
    sgm_pending = int(pending_by_module.get("sgm", 0))

    # Expiring docs (docs library)
    # next 30 days window
    exp_to = d_to + datetime.timedelta(days=30)
    exp_clause, exp_params = _station_filter_clause(station_id)
    cur.execute(
        f"""
        SELECT COUNT(1) AS c
        FROM document_versions
        WHERE brand=? AND expires_at IS NOT NULL
          AND date(expires_at) BETWEEN date(?) AND date(?)
          {exp_clause.replace('station_id', 'station_id')}
        """,
        (brand, d_to.isoformat(), exp_to.isoformat(), *exp_params),
    )
    docs_expiring_30d = int(cur.fetchone()["c"] or 0)

    # Station ranking snapshot (admin dashboard)
    cur = conn.cursor()
    cur.execute("SELECT id, code, name FROM stations WHERE brand=? ORDER BY code ASC, id ASC", (brand,))
    ranking = []
    for st_row in cur.fetchall():
        sid = int(st_row["id"])
        cur.execute("SELECT COUNT(1) AS c FROM alerts WHERE brand=? AND station_id=? AND status='open'", (brand, sid))
        a_open = int(cur.fetchone()["c"] or 0)
        cur.execute("SELECT COUNT(1) AS c FROM payments WHERE brand=? AND station_id=? AND status='pending'", (brand, sid))
        p_pending = int(cur.fetchone()["c"] or 0)
        cur.execute("SELECT COUNT(1) AS c FROM doc_submissions ds JOIN doc_requirements r ON r.id=ds.requirement_id AND r.brand=ds.brand AND r.module=ds.module WHERE ds.brand=? AND r.station_id=? AND ds.review_status='PENDING'", (brand, sid))
        doc_pending = int(cur.fetchone()["c"] or 0)
        score = max(0, 100 - (a_open * 8) - (p_pending * 10) - (doc_pending * 12))
        ranking.append({"station_id": sid, "code": st_row["code"], "name": st_row["name"], "alerts_open": a_open, "payments_pending": p_pending, "doc_pending": doc_pending, "score": score})
    ranking.sort(key=lambda x: (-int(x.get("score") or 0), x.get("code") or ""))

    conn.close()

    return {
        "brand": brand,
        "station": st,
        "range": {"from": d_from.isoformat(), "to": d_to.isoformat()},
        "ranking": ranking[:10],
        "metrics": {
            "alerts_open": alerts_open,
            "alerts_created": alerts_created,
            "maintenance_created": maint_created,
            "pipas_created": pipas_created,
            "payments_pending": payments_pending,
            "payments_validated": payments_validated,
            "events_planned": events_planned,
            "submissions": {
                "submitted": sub_by_status.get("submitted", 0),
                "reviewed": sub_by_status.get("reviewed", 0),
                "approved": sub_by_status.get("approved", 0),
                "rejected": sub_by_status.get("rejected", 0),
            },
            "sasisopa_pending_review": sasisopa_pending,
            "sgm_pending_review": sgm_pending,
            "docs_expiring_30d": docs_expiring_30d,
        },
    }


def _chart_categories(metrics: dict) -> list[tuple[str, int]]:
    """Return (label, value) pairs for executive bar charts."""
    sub = metrics.get("submissions") or {}
    return [
        ("Alertas (rango)", int(metrics.get("alerts_created", 0) or 0)),
        ("Mantenimientos (rango)", int(metrics.get("maintenance_created", 0) or 0)),
        ("Pipas (rango)", int(metrics.get("pipas_created", 0) or 0)),
        ("Eventos (rango)", int(metrics.get("events_planned", 0) or 0)),
        ("Aprobadas (rango)", int(sub.get("approved", 0) or 0)),
        ("Rechazadas (rango)", int(sub.get("rejected", 0) or 0)),
        ("Pagos pendientes", int(metrics.get("payments_pending", 0) or 0)),
        ("Docs por vencer (30d)", int(metrics.get("docs_expiring_30d", 0) or 0)),
        ("SASISOPA por revisar", int(metrics.get("sasisopa_pending_review", 0) or 0)),
        ("SGM por revisar", int(metrics.get("sgm_pending_review", 0) or 0)),
    ]


def register(app):
    ctx = app.extensions["ctx"]
    login_required = ctx.login_required
    role_required = ctx.role_required

    @app.get("/admin/executive")
    @login_required
    @role_required("admin")
    def admin_executive_page():
        return render_template("admin/executive.html", me=ctx.get_me())

    @app.get("/api/admin/executive/summary")
    @login_required
    @role_required("admin")
    def api_admin_executive_summary():
        brand = get_brand()
        station_id = request.args.get("station_id")
        sid = None
        if station_id:
            try:
                sid = int(station_id)
            except Exception:
                return jsonify({"ok": False, "error": "invalid_station_id"}), 400
        d_from, d_to = _date_range_from_request()
        data = _compute_summary(brand, sid, d_from, d_to)
        ctx.log_action(ctx.get_me(), "view_executive_summary", "reports", f"{sid or 'all'}:{d_from.isoformat()}:{d_to.isoformat()}")
        return jsonify({"ok": True, **data})

    @app.get("/api/admin/executive/export.xlsx")
    @login_required
    @role_required("admin")
    def api_admin_executive_export_xlsx():
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        import tempfile

        brand = get_brand()
        station_id = request.args.get("station_id")
        sid = None
        if station_id:
            try:
                sid = int(station_id)
            except Exception:
                sid = None
        d_from, d_to = _date_range_from_request()
        data = _compute_summary(brand, sid, d_from, d_to)

        wb = Workbook()
        ws = wb.active
        ws.title = "Resumen"
        ws.append(["Reporte ejecutivo", brand])
        ws.append(["Rango", f"{data['range']['from']} a {data['range']['to']}"])
        if data.get("station"):
            ws.append(["Estación", f"{data['station']['name']} ({data['station']['code']})"])
        else:
            ws.append(["Estación", "Todas"])
        ws.append([])

        m = data["metrics"]
        ws.append(["Métrica", "Valor"])
        ws.append(["Alertas abiertas", m["alerts_open"]])
        ws.append(["Alertas creadas (rango)", m["alerts_created"]])
        ws.append(["Mantenimientos (rango)", m["maintenance_created"]])
        ws.append(["Pipas (rango)", m["pipas_created"]])
        ws.append(["Mensualidades pendientes", m["payments_pending"]])
        ws.append(["Mensualidades validadas", m["payments_validated"]])
        ws.append(["Eventos programados (rango)", m["events_planned"]])
        ws.append(["Entregas aprobadas (rango)", m["submissions"]["approved"]])
        ws.append(["Entregas rechazadas (rango)", m["submissions"]["rejected"]])
        ws.append(["SASISOPA por revisar", m["sasisopa_pending_review"]])
        ws.append(["SGM por revisar", m["sgm_pending_review"]])
        ws.append(["Docs por vencer (30 días)", m["docs_expiring_30d"]])

        # Bars sheet (with embedded chart)
        ws2 = wb.create_sheet("Barras")
        ws2.append(["Categoría", "Valor"])
        cats = _chart_categories(m)
        for label, val in cats:
            ws2.append([label, int(val)])

        chart = BarChart()
        chart.type = "col"
        chart.title = "Resumen (barras)"
        chart.y_axis.title = "Valor"
        chart.x_axis.title = "Categoría"
        data_ref = Reference(ws2, min_col=2, min_row=1, max_row=1 + len(cats))
        cat_ref = Reference(ws2, min_col=1, min_row=2, max_row=1 + len(cats))
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cat_ref)
        chart.height = 10
        chart.width = 22
        ws2.add_chart(chart, "D2")

        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        wb.save(tmp.name)
        tmp.close()

        ctx.log_action(ctx.get_me(), "download_executive_xlsx", "reports", f"{sid or 'all'}:{d_from.isoformat()}:{d_to.isoformat()}")
        return send_file(tmp.name, as_attachment=True, download_name=f"reporte_ejecutivo_{brand}_{sid or 'all'}_{d_from.isoformat()}_{d_to.isoformat()}.xlsx")

    @app.get("/api/admin/executive/export.pdf")
    @login_required
    @role_required("admin")
    def api_admin_executive_export_pdf():
        """PDF executive report with a simple bar chart."""
        import tempfile
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas

        brand = get_brand()
        station_id = request.args.get("station_id")
        sid = None
        if station_id:
            try:
                sid = int(station_id)
            except Exception:
                sid = None
        d_from, d_to = _date_range_from_request()
        data = _compute_summary(brand, sid, d_from, d_to)
        me = ctx.get_me()

        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
        c = canvas.Canvas(tmp.name, pagesize=letter)
        w, h = letter

        # Header
        c.setFont("Helvetica-Bold", 15)
        c.drawString(40, h - 50, "COG WORK LOG - Reporte ejecutivo")
        c.setFont("Helvetica", 10)
        st_line = "Estación: Todas"
        if data.get("station"):
            st_line = f"Estación: {data['station']['name']} ({data['station']['code']})"
        c.drawString(40, h - 70, st_line)
        c.drawString(40, h - 84, f"Rango: {data['range']['from']} a {data['range']['to']}")
        c.drawString(40, h - 98, f"Generado por: {me.get('username','-')} • {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}")

        # Table of metrics
        m = data["metrics"]
        rows = [
            ("Alertas abiertas", m["alerts_open"]),
            ("Alertas creadas (rango)", m["alerts_created"]),
            ("Mantenimientos (rango)", m["maintenance_created"]),
            ("Pipas (rango)", m["pipas_created"]),
            ("Mensualidades pendientes", m["payments_pending"]),
            ("Mensualidades validadas", m["payments_validated"]),
            ("Eventos programados (rango)", m["events_planned"]),
            ("Entregas aprobadas (rango)", m["submissions"]["approved"]),
            ("Entregas rechazadas (rango)", m["submissions"]["rejected"]),
            ("SASISOPA por revisar", m["sasisopa_pending_review"]),
            ("SGM por revisar", m["sgm_pending_review"]),
            ("Docs por vencer (30 días)", m["docs_expiring_30d"]),
        ]

        y = h - 130
        c.setFont("Helvetica-Bold", 10)
        c.drawString(40, y, "Métricas")
        y -= 14
        c.setFont("Helvetica-Bold", 9)
        c.drawString(40, y, "Concepto")
        c.drawString(330, y, "Valor")
        y -= 6
        c.line(40, y, w - 40, y)
        y -= 12
        c.setFont("Helvetica", 9)
        for name, val in rows:
            if y < 360:
                break
            c.drawString(40, y, str(name))
            c.drawRightString(w - 40, y, f"{int(val)}")
            y -= 12

        # Simple bar chart
        cats = _chart_categories(m)
        chart_top = 330
        chart_left = 40
        chart_w = w - 80
        chart_h = 220

        # Frame
        c.setFont("Helvetica-Bold", 10)
        c.drawString(chart_left, chart_top + chart_h + 10, "Barras (resumen)")
        c.rect(chart_left, chart_top, chart_w, chart_h)

        values = [int(v) for _, v in cats]
        max_v = max(values) if values else 1
        if max_v <= 0:
            max_v = 1

        n = len(cats) if cats else 1
        gap = 6
        bar_w = max(10, (chart_w - gap * (n + 1)) / n)
        x = chart_left + gap
        c.setFont("Helvetica", 7)
        for (label, val) in cats:
            v = int(val)
            bh = (v / max_v) * (chart_h - 26)
            # bar
            c.rect(x, chart_top + 18, bar_w, bh, fill=0)
            # value
            c.drawCentredString(x + bar_w / 2, chart_top + 18 + bh + 2, str(v))
            # label (short)
            short = label
            if len(short) > 16:
                short = short[:16] + "…"
            c.drawCentredString(x + bar_w / 2, chart_top + 4, short)
            x += bar_w + gap

        c.showPage()
        c.save()
        tmp.close()

        ctx.log_action(me, "download_executive_pdf", "reports", f"{sid or 'all'}:{d_from.isoformat()}:{d_to.isoformat()}")
        return send_file(
            tmp.name,
            as_attachment=True,
            download_name=f"reporte_ejecutivo_{brand}_{sid or 'all'}_{d_from.isoformat()}_{d_to.isoformat()}.pdf",
        )
